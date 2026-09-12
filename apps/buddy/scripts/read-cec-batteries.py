"""Read-only extraction. Preserve physical worksheet rows and reject formulas."""
import datetime
import json
import sys
import openpyxl

if openpyxl.__version__ != "3.1.5":
    raise RuntimeError("This adapter requires openpyxl==3.1.5")
workbook = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=False, keep_links=False)
try:
    if workbook.sheetnames != ["Battery"]:
        raise ValueError("CEC worksheet changed")
    sheet = workbook["Battery"]
    if sheet.max_row > 10000 or sheet.max_column != 16:
        raise ValueError("Unexpected CEC workbook dimensions")
    rows = []
    for row in sheet:
        if any(cell.data_type == "f" for cell in row):
            raise ValueError("Unexpected formula; do not substitute cached results")
        rows.append([cell.value.isoformat()[:10] if isinstance(cell.value, (datetime.datetime, datetime.date)) else cell.value for cell in row])
    print(json.dumps({"sheet": sheet.title, "rows": rows}, ensure_ascii=False))
finally:
    workbook.close()
